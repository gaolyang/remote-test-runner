from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.testcase.models import TestCase


class ExcelImportError(ValueError):
    pass


HEADER_ALIASES = {
    "case_id": {"测试编号", "用例编号", "测试id", "用例id", "caseid", "testcaseid"},
    "name": {"测试名称", "用例名称", "测试功能", "功能名称", "testname", "testcasename"},
    "steps": {"测试执行步骤", "执行步骤", "测试步骤", "步骤", "命令", "command", "steps"},
    "step_name": {"步骤名称", "stepname"},
    "expected": {"预期结果", "预期输出", "expected", "expectedresult"},
    "role": {"步骤类型", "角色", "role"},
    "interactive": {"是否交互", "交互", "interactive"},
    "description": {"描述", "测试描述", "用例描述", "description"},
}


def _normalise_header(value: Any) -> str:
    return re.sub(r"[\s_()（）\-]+", "", str(value or "")).lower()


def _header_map(header: list[Any]) -> dict[str, int]:
    normalised = [_normalise_header(value) for value in header]
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        alias_set = {_normalise_header(item) for item in aliases}
        for index, value in enumerate(normalised):
            if value in alias_set:
                result[field] = index
                break
    missing = [label for field, label in (("name", "测试名称"), ("steps", "测试执行步骤")) if field not in result]
    if missing:
        raise ExcelImportError(f"缺少必需列：{', '.join(missing)}")
    return result


def _text(row: list[Any], columns: dict[str, int], field: str) -> str:
    index = columns.get(field)
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _split_description_and_command(value: str) -> tuple[str | None, str]:
    """Split a Chinese step label from its shell command conservatively.

    Supported examples use either a full-width or ASCII colon.  Requiring the
    label to start with Chinese text avoids breaking valid commands such as
    ``echo key:value`` and URLs containing ``https:``.
    """

    match = re.match(r"^(?P<description>[\u4e00-\u9fff（(【\[].*?)[：:]\s*(?P<command>.+)$", value)
    if match is None:
        return None, value
    description = match.group("description").strip()
    command = match.group("command").strip()
    if not re.match(r"^(?:[A-Za-z_][A-Za-z0-9_.+-]*|[./~$])", command):
        return None, value
    return description, command


def _split_steps(value: str) -> list[tuple[str | None, str]]:
    lines: list[tuple[str | None, str]] = []
    for raw in value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        step_text = re.sub(r"^\s*(?:步骤\s*)?\d+\s*[.、):：-]\s*", "", raw).strip()
        step_text = re.sub(r"^\s*[-•]\s*", "", step_text).strip()
        if step_text:
            lines.append(_split_description_and_command(step_text))
    return lines


def _truthy(value: str) -> bool:
    return value.strip().lower() in {"1", "true", "yes", "y", "是", "需要", "交互"}


def _role(value: str) -> str:
    mapping = {"准备": "setup", "前置": "setup", "验证": "verify", "检查": "verify", "清理": "cleanup", "后置": "cleanup"}
    role = mapping.get(value.strip().lower(), value.strip().lower() or "verify")
    if role not in {"setup", "verify", "cleanup"}:
        raise ExcelImportError(f"无法识别的步骤类型：{value}")
    return role


def _expected(value: str) -> dict[str, Any]:
    expected: dict[str, Any] = {"exit_code": 0}
    if not value:
        return expected
    match = re.search(r"(?:exit[_ ]?code|退出码)\s*[=:：]\s*(-?\d+)", value, re.IGNORECASE)
    if match:
        expected["exit_code"] = int(match.group(1))
        value = (value[: match.start()] + value[match.end() :]).strip(" ;；,")
    if value:
        expected["stdout_contains"] = [item.strip() for item in re.split(r"[\n；;]+", value) if item.strip()]
    return expected


def _make_case_id(raw: str, index: int) -> str:
    candidate = re.sub(r"[^A-Za-z0-9_.-]+", "_", raw.strip()).strip("_.-")
    return candidate or f"TC_IMPORT_{index:03d}"


def parse_tabular_testcases(rows: list[list[Any]]) -> list[TestCase]:
    if not rows:
        raise ExcelImportError("文件中没有数据")
    columns = _header_map(rows[0])
    groups: OrderedDict[str, dict[str, Any]] = OrderedDict()
    current_name = ""
    for row_number, row in enumerate(rows[1:], start=2):
        name = _text(row, columns, "name") or current_name
        step_text = _text(row, columns, "steps")
        if not name and not step_text:
            continue
        if not name:
            raise ExcelImportError(f"第 {row_number} 行缺少测试名称")
        current_name = name
        if not step_text:
            continue
        group = groups.setdefault(
            name,
            {
                "raw_id": _text(row, columns, "case_id"),
                "description": _text(row, columns, "description"),
                "steps": [],
            },
        )
        parsed_steps = _split_steps(step_text)
        if not parsed_steps:
            raise ExcelImportError(f"第 {row_number} 行没有可执行步骤")
        role = _role(_text(row, columns, "role"))
        expected = _expected(_text(row, columns, "expected"))
        interactive = _truthy(_text(row, columns, "interactive"))
        supplied_step_name = _text(row, columns, "step_name")
        for parsed_name, command in parsed_steps:
            group["steps"].append(
                {
                    "role": role,
                    "name": supplied_step_name or parsed_name or command,
                    "action": {"type": "shell", "command": command},
                    "expected": expected,
                    "interaction": {"mode": "manual"} if interactive else None,
                }
            )
    if not groups:
        raise ExcelImportError("没有找到可导入的测试步骤")

    cases: list[TestCase] = []
    used_ids: set[str] = set()
    for index, (name, group) in enumerate(groups.items(), start=1):
        case_id = _make_case_id(group["raw_id"], index)
        if case_id in used_ids:
            raise ExcelImportError(f"测试编号重复：{case_id}")
        used_ids.add(case_id)
        steps = []
        for step_index, item in enumerate(group["steps"], start=1):
            steps.append({"id": step_index, **item})
        cases.append(
            TestCase.model_validate(
                {
                    "version": "1.0",
                    "testcase": {"id": case_id, "name": name, "description": group["description"]},
                    "target": {"connection": {"type": "ssh"}},
                    "steps": steps,
                }
            )
        )
    return cases


def parse_testcase_upload(filename: str, content: bytes) -> list[TestCase]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".xlsx", ".xlsm"}:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
        elif suffix == ".csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gb18030")
            rows = [list(row) for row in csv.reader(io.StringIO(text))]
        else:
            raise ExcelImportError("仅支持 .xlsx、.xlsm 或 .csv 文件")
    except ExcelImportError:
        raise
    except Exception as exc:
        raise ExcelImportError(f"无法读取表格：{exc}") from exc
    return parse_tabular_testcases(rows)
